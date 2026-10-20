"""DRBench (https://github.com/ServiceNow/drbench) in local-file (L1) mode.

Each task is an enterprise deep-research question grounded in a fixed set of
company files: PDFs, Word documents, slide decks, spreadsheets, mailbox
exports, and Mattermost chat exports. ``env.json`` labels every file as an
``insight`` or ``distractor`` source; ``eval.json`` holds the gold insights.

L1 mode replaces the official Docker services with a plain workspace:

* only files listed in ``env.json`` are exposed, at their ``destination``
  path, as deterministic text renderings (``shared/<name>.<ext>.md``);
* answer material (``eval.json``, ``qa_dict.json``, ``file_dict.json``) never
  enters the workspace;
* the Agent cites original file names, matching the official citation format.

Renderings are cached under ``<benchmark dir>/_derived/<task_id>/`` and keyed
by the source SHA-256 and :data:`RENDERER_VERSION`.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
from collections.abc import Iterable
from datetime import UTC, datetime
from pathlib import Path, PurePosixPath
from typing import Any

from benchmarks.public.reports.models import TASK_KIND_REPORT, BenchmarkTask

NAME = "drbench"
DIRECTORY_CANDIDATES = ("drbench", "DRBench", "DrBench")
SUBSETS = ("val", "minival", "sanity")
DIFFICULTIES = ("easy", "medium", "hard")
RENDERER_VERSION = 2  # FrontierAgent uses its existing pypdf dependency for PDFs.
DERIVED_DIRNAME = "_derived"
FILE_INDEX_NAME = "FILES.md"
RENDER_SUFFIX = ".md"
_SOURCE_PREFIX = "drbench/data/"
_FORBIDDEN_WORKSPACE_NAMES = {"eval.json", "qa_dict.json", "file_dict.json"}
_MAX_SHEET_ROWS = 2_000
_APP_LABELS = {
    "nextcloud": "Nextcloud file",
    "file_system": "File system file",
    "email": "Email mailbox export (RoundCube)",
    "mattermost": "Mattermost chat export",
}


class DRBenchDataError(ValueError):
    """Raised when DRBench data is missing or cannot be rendered."""


# ---------------------------------------------------------------------------
# Layout and loading
# ---------------------------------------------------------------------------


def find_benchmark_dir(data_root: Path) -> Path:
    for directory in DIRECTORY_CANDIDATES:
        candidate = data_root / directory
        if candidate.is_dir():
            return candidate
    searched = ", ".join(str(data_root / value) for value in DIRECTORY_CANDIDATES)
    raise DRBenchDataError(
        f"DRBench data directory was not found. Searched: {searched}. "
        f"Run `python -m benchmarks.public.runner.run_reports prepare {NAME}`."
    )


def resolve_data_dir(benchmark_dir: Path) -> Path:
    """Accept both ``<dir>/{subsets,tasks}`` and a full repository clone."""

    for candidate in (benchmark_dir, benchmark_dir / "drbench" / "data"):
        if (candidate / "tasks").is_dir() and (candidate / "subsets").is_dir():
            return candidate
    raise DRBenchDataError(
        f"{benchmark_dir} must contain `subsets/` and `tasks/` (the contents of "
        "the drbench repository's `drbench/data/`), or be a full repository clone."
    )


def load_tasks(
    data_root: Path,
    *,
    data_file: Path | None = None,
    subset: str = "val",
    difficulty: str | None = None,
) -> list[BenchmarkTask]:
    if data_file is not None:
        benchmark_dir = Path(data_file).expanduser().resolve()
        if benchmark_dir.is_file():
            raise DRBenchDataError("--data-file PATH must point to the DRBench data directory.")
    else:
        benchmark_dir = find_benchmark_dir(data_root)
    data_dir = resolve_data_dir(benchmark_dir)
    if subset not in SUBSETS:
        raise DRBenchDataError(f"--drbench-subset must be one of {', '.join(SUBSETS)}.")
    if difficulty is not None and difficulty not in DIFFICULTIES:
        raise DRBenchDataError(f"--drbench-difficulty must be one of {', '.join(DIFFICULTIES)}.")

    task_ids = _subset_task_ids(data_dir, subset)
    tasks: list[BenchmarkTask] = []
    for task_id in task_ids:
        task = load_task(data_dir, task_id, derived_root=benchmark_dir / DERIVED_DIRNAME)
        if task is None:
            continue
        if difficulty is not None and task.metadata.get("difficulty") != difficulty:
            continue
        tasks.append(task)
    if not tasks:
        raise DRBenchDataError(f"No DRBench tasks remain in subset {subset!r} after filters.")
    return tasks


def _subset_task_ids(data_dir: Path, subset: str) -> list[str]:
    path = data_dir / "subsets" / f"{subset}.jsonl"
    if not path.is_file():
        raise DRBenchDataError(f"DRBench subset file does not exist: {path}")
    task_ids: list[str] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        row = json.loads(line)
        if isinstance(row, dict) and row.get("task_id"):
            task_ids.append(str(row["task_id"]))
    if not task_ids and subset == "sanity":
        task_ids = ["SANITY0"]
    return task_ids


def load_task(data_dir: Path, task_id: str, *, derived_root: Path) -> BenchmarkTask | None:
    if not re.fullmatch(r"[A-Za-z0-9_-]+", task_id):
        raise DRBenchDataError(f"Unsafe DRBench task ID: {task_id!r}")
    task_dir = data_dir / "tasks" / task_id
    config_dir = task_dir / "config"
    if not (config_dir / "task.json").is_file() or not (config_dir / "env.json").is_file():
        # DR0000 ships only info.json; it is not a runnable task.
        return None
    task_config = _read_json(config_dir / "task.json")
    env_config = _read_json(config_dir / "env.json")
    eval_config = (
        _read_json(config_dir / "eval.json") if (config_dir / "eval.json").is_file() else {}
    )
    info = _read_json(task_dir / "info.json") if (task_dir / "info.json").is_file() else {}

    files: list[dict[str, str]] = []
    for entry in env_config.get("env_files") or []:
        if not isinstance(entry, dict):
            continue
        source = _resolve_source(data_dir, str(entry.get("source") or ""))
        destination = _safe_destination(str(entry.get("destination") or ""))
        if source.name.casefold() in _FORBIDDEN_WORKSPACE_NAMES:
            raise DRBenchDataError("An answer/evaluation file cannot be a task input.")
        if not source.is_file():
            raise DRBenchDataError(f"DRBench {task_id}: source file is missing: {source}")
        files.append(
            {
                "source": str(source),
                "destination": destination,
                "app": str(entry.get("app") or "file_system"),
                "qa_type": str(entry.get("qa_type") or ""),
            }
        )
    qa = [
        {
            key: item.get(key)
            for key in (
                "id",
                "type",
                "qa_type",
                "question",
                "answer",
                "supporting_file_paths",
                "supporting_urls",
            )
            if key in item
        }
        for item in eval_config.get("dr_report_evaluation_qa") or []
        if isinstance(item, dict)
    ]
    question = str(task_config.get("dr_question") or "").strip()
    if not question:
        raise DRBenchDataError(f"DRBench {task_id}: task.json has no dr_question.")
    persona = {
        key: value
        for key, value in (task_config.get("persona") or {}).items()
        if key not in {"password"}
    }
    company = task_config.get("company_info") or {}
    insight_qa = [item for item in qa if item.get("qa_type") == "insight"]
    return BenchmarkTask(
        benchmark=NAME,
        task_id=str(task_config.get("task_id") or task_id),
        question=question,
        answers=(),
        kind=TASK_KIND_REPORT,
        reference={"files": files, "qa": qa},
        metadata={
            "difficulty": str(info.get("difficulty") or ""),
            "industry": str(info.get("industry") or ""),
            "domain": str(info.get("domain") or ""),
            "company": str(company.get("name") or ""),
            "date": str(task_config.get("date") or ""),
            "n_files": len(files),
            "n_insight_qa": len(insight_qa),
            "n_external_insight_qa": sum(
                item.get("type") == "external_fact" for item in insight_qa
            ),
            "data_dir": str(data_dir),
            "derived_dir": str(derived_root / task_id),
            # Prompt context only; no labels or answers.
            "context": {
                "company_info": company,
                "persona": persona,
                "date": str(task_config.get("date") or ""),
            },
            "workspace_files": [
                {key: item[key] for key in ("source", "destination", "app")} for item in files
            ],
        },
    )


def _read_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DRBenchDataError(f"Cannot read {path}: {exc}") from exc


def _resolve_source(data_dir: Path, source: str) -> Path:
    normalized = source.replace("\\", "/")
    if normalized.startswith(_SOURCE_PREFIX):
        path = data_dir / normalized[len(_SOURCE_PREFIX) :]
    elif (marker := normalized.find("/tasks/")) >= 0:
        path = data_dir / normalized[marker + 1 :]
    else:
        path = Path(normalized)
        path = path if path.is_absolute() else data_dir / normalized
    path = path.resolve()
    if not path.is_relative_to(data_dir.resolve()):
        raise DRBenchDataError("DRBench input sources must stay inside the dataset directory.")
    if path.is_file() or not path.parent.is_dir():
        return path
    # Upstream data was authored on a case-insensitive filesystem; at least
    # one env.json entry (DR0038) differs from the file on disk only by case.
    folded = path.name.casefold()
    matches = [item for item in path.parent.iterdir() if item.name.casefold() == folded]
    return matches[0] if len(matches) == 1 else path


def _safe_destination(destination: str) -> str:
    path = PurePosixPath(destination.replace("\\", "/").lstrip("/"))
    if (
        not path.parts
        or any(part in {"..", ""} for part in path.parts)
        or path.name.casefold() in _FORBIDDEN_WORKSPACE_NAMES
    ):
        raise DRBenchDataError(f"Unsafe DRBench destination path: {destination!r}")
    return str(path)


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def ensure_rendered(task: BenchmarkTask) -> dict[str, Any]:
    """Render every workspace file of ``task`` once; return the cache index."""

    derived_dir = Path(task.metadata["derived_dir"])
    index_path = derived_dir / "index.json"
    files = task.metadata["workspace_files"]
    expected = {item["destination"]: _sha256(Path(item["source"])) for item in files}
    if index_path.is_file():
        try:
            index = json.loads(index_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            index = {}
        cached = {item["destination"]: item for item in index.get("files", [])}
        if (
            index.get("renderer_version") == RENDERER_VERSION
            and set(cached) == set(expected)
            and all(cached[key].get("source_sha256") == value for key, value in expected.items())
            and all((derived_dir / cached[key]["rendered"]).is_file() for key in cached)
        ):
            return index

    if derived_dir.exists():
        shutil.rmtree(derived_dir)
    derived_dir.mkdir(parents=True, exist_ok=True)
    rendered_files = []
    for item in files:
        source = Path(item["source"])
        text, keys = render_file(source, app=item["app"], destination=item["destination"])
        rendered = item["destination"] + RENDER_SUFFIX
        target = derived_dir / rendered
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
        rendered_files.append(
            {
                "destination": item["destination"],
                "rendered": rendered,
                "app": item["app"],
                "source_sha256": expected[item["destination"]],
                "chars": len(text),
                "citation_keys": keys,
            }
        )
    index = {
        "task_id": task.task_id,
        "renderer_version": RENDERER_VERSION,
        "files": rendered_files,
    }
    index_path.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    return index


def render_file(source: Path, *, app: str, destination: str) -> tuple[str, list[str]]:
    """Return ``(markdown text, citation keys)`` for one DRBench source file."""

    suffix = source.suffix.lower()
    name = PurePosixPath(destination).name
    keys = _name_keys(name)
    if app == "email" or (suffix == ".jsonl" and "email" in source.name.lower()):
        body, extra = _render_mailbox(source, name)
    elif app == "mattermost" or (suffix == ".jsonl" and "mattermost" in source.name.lower()):
        body, extra = _render_mattermost(source, name)
    elif suffix == ".pdf":
        body, extra = _render_pdf(source), []
    elif suffix == ".docx":
        body, extra = _render_docx(source), []
    elif suffix == ".pptx":
        body, extra = _render_pptx(source), []
    elif suffix in {".xlsx", ".xlsm"}:
        body, extra = _render_xlsx(source), []
    elif suffix in {".csv", ".md", ".txt", ".json", ".jsonl"}:
        body, extra = source.read_text(encoding="utf-8", errors="replace"), []
    else:
        raise DRBenchDataError(f"Unsupported DRBench file type: {source}")
    header = (
        f"<!-- Text rendering of `{destination}` ({_APP_LABELS.get(app, app)}). "
        f"Cite it as `{name}`. -->\n\n"
    )
    return header + body.strip() + "\n", _dedupe(keys + extra)


def _missing_dependency(package: str) -> DRBenchDataError:
    return DRBenchDataError(
        f"Rendering DRBench files requires `{package}`. "
        "Install the benchmark extras: `uv sync --extra eval`."
    )


def _render_pdf(source: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise _missing_dependency("pypdf") from exc
    parts = []
    for number, page in enumerate(PdfReader(str(source)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(f"## Page {number}\n\n{text}")
    return "\n\n".join(parts)


def _render_docx(source: Path) -> str:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise _missing_dependency("python-docx") from exc
    document = docx.Document(str(source))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = re.match(r"Heading (\d)", style)
            if level:
                parts.append(f"{'#' * min(6, int(level.group(1)) + 1)} {text}")
            elif style == "Title":
                parts.append(f"# {text}")
            elif "List" in style:
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            parts.append(_markdown_table(rows))
    return "\n\n".join(part for part in parts if part)


def _render_pptx(source: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise _missing_dependency("python-pptx") from exc
    presentation = Presentation(str(source))
    slides: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts = [f"## Slide {number}"]
        for shape in slide.shapes:
            parts.extend(_pptx_shape_text(shape))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Speaker notes: {notes}")
        slides.append("\n\n".join(parts))
    return "\n\n".join(slides)


def _pptx_shape_text(shape: Any) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        text = "\n".join(
            paragraph.text.strip()
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        )
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False) and shape.has_table:
        rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
        parts.append(_markdown_table(rows))
    if getattr(shape, "has_chart", False) and shape.has_chart:
        chart = shape.chart
        lines = ["Chart data:"]
        try:
            categories = [str(value) for value in chart.plots[0].categories]
        except (IndexError, AttributeError):
            categories = []
        for plot in chart.plots:
            for series in plot.series:
                values = ", ".join(
                    f"{categories[i] if i < len(categories) else i}: {value}"
                    for i, value in enumerate(series.values)
                )
                lines.append(f"- {series.name}: {values}")
        parts.append("\n".join(lines))
    for child in getattr(shape, "shapes", []) or []:
        parts.extend(_pptx_shape_text(child))
    return parts


def _render_xlsx(source: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise _missing_dependency("openpyxl") from exc
    workbook = openpyxl.load_workbook(str(source), data_only=True, read_only=True)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            truncated = False
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    continue
                if len(rows) >= _MAX_SHEET_ROWS:
                    truncated = True
                    break
                rows.append(values)
            section = [f"## Sheet: {sheet.title}", _markdown_table(rows) or "(empty)"]
            if truncated:
                section.append(f"(truncated after {_MAX_SHEET_ROWS} rows)")
            parts.append("\n\n".join(section))
    finally:
        workbook.close()
    return "\n\n".join(parts)


def _render_mailbox(source: Path, name: str) -> tuple[str, list[str]]:
    users: list[str] = []
    emails: list[dict[str, Any]] = []
    for record in _jsonl(source):
        if record.get("type") == "user":
            full_name = " ".join(
                str(record.get(key) or "") for key in ("first_name", "last_name")
            ).strip()
            address = str(record.get("email") or "")
            if address:
                users.append(f"- {full_name or record.get('username', '')} <{address}>")
        elif record.get("type") == "email":
            emails.append(record)
    emails.sort(key=lambda item: str(item.get("date") or ""))
    parts = [f"# Mailbox export: {name}"]
    if users:
        parts.append("## Address book\n\n" + "\n".join(users))
    keys: list[str] = []
    for item in emails:
        subject = str(item.get("subject") or "(no subject)").strip()
        if len(subject) >= 12:
            keys.append(subject.casefold())
        sender = str(item.get("from") or "")
        sender_name = str(item.get("from_name") or "").strip()
        lines = [
            f"## Email: {subject}",
            f"From: {sender_name} <{sender}>" if sender_name else f"From: {sender}",
            f"To: {', '.join(_strings(item.get('to')))}",
        ]
        cc = _strings(item.get("cc"))
        if cc:
            lines.append(f"Cc: {', '.join(cc)}")
        lines.append(f"Date: {item.get('date') or ''}")
        attachments = [
            str(value.get("filename") or value.get("name") or value)
            if isinstance(value, dict)
            else str(value)
            for value in item.get("attachments") or []
        ]
        if attachments:
            lines.append(f"Attachments: {', '.join(attachments)}")
        lines.append("")
        lines.append(str(item.get("body") or "").strip())
        parts.append("\n".join(lines))
    return "\n\n".join(parts), keys


def _render_mattermost(source: Path, name: str) -> tuple[str, list[str]]:
    channels: dict[tuple[str, str], str] = {}
    posts: list[dict[str, Any]] = []
    users: list[str] = []
    for record in _jsonl(source):
        kind = record.get("type")
        if kind == "channel" and isinstance(record.get("channel"), dict):
            channel = record["channel"]
            channels[(str(channel.get("team") or ""), str(channel.get("name") or ""))] = str(
                channel.get("display_name") or channel.get("name") or ""
            )
        elif kind == "user" and isinstance(record.get("user"), dict):
            user = record["user"]
            label = " ".join(
                str(user.get(key) or "") for key in ("first_name", "last_name")
            ).strip()
            position = str(user.get("position") or "").strip()
            users.append(
                f"- {user.get('username', '')}: {label}" + (f", {position}" if position else "")
            )
        elif kind == "post" and isinstance(record.get("post"), dict):
            posts.append(record["post"])
    posts.sort(key=lambda post: int(post.get("create_at") or 0))
    parts = [f"# Mattermost chat export: {name}"]
    if users:
        parts.append("## Members\n\n" + "\n".join(users))
    keys: list[str] = []
    by_channel: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for post in posts:
        key = (str(post.get("team") or ""), str(post.get("channel") or ""))
        by_channel.setdefault(key, []).append(post)
    for (team, channel), channel_posts in by_channel.items():
        display = channels.get((team, channel), channel)
        for value in (channel, display):
            if len(value) >= 6:
                keys.append(value.casefold())
        lines = [f"## Team: {team} / Channel: {channel} ({display})"]
        for post in channel_posts:
            lines.append(
                f"[{_timestamp(post.get('create_at'))}] "
                f"(User: {post.get('user', '')}, Team: {team}, Channel: {channel}) "
                f"{str(post.get('message') or '').strip()}"
            )
        parts.append("\n".join(lines))
    return "\n\n".join(parts), keys


def _jsonl(source: Path) -> Iterable[dict[str, Any]]:
    for line in source.read_text(encoding="utf-8", errors="replace").splitlines():
        if not line.strip():
            continue
        try:
            record = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(record, dict):
            yield record


def _timestamp(value: Any) -> str:
    try:
        moment = datetime.fromtimestamp(int(value) / 1000, tz=UTC)
    except (TypeError, ValueError, OverflowError, OSError):
        return ""
    return moment.strftime("%Y-%m-%d %H:%M UTC")


def _strings(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    return [str(value)] if value else []


def _markdown_table(rows: list[list[str]]) -> str:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)

    def cell(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", " ").strip()

    lines = []
    for index, row in enumerate(rows):
        padded = [cell(value) for value in row] + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded) + " |")
        if index == 0:
            lines.append("|" + "---|" * width)
    return "\n".join(lines)


def _name_keys(name: str) -> list[str]:
    folded = name.casefold()
    keys = [folded]
    stem = PurePosixPath(folded).stem
    if len(stem) >= 8 and stem != folded:
        keys.append(stem)
    return keys


def _dedupe(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ---------------------------------------------------------------------------
# Workspace
# ---------------------------------------------------------------------------


def stage_workspace(task: BenchmarkTask, workspace: Path) -> dict[str, str]:
    """Copy rendered files into ``workspace``; return rendered→original paths."""

    index = ensure_rendered(task)
    derived_dir = Path(task.metadata["derived_dir"])
    mapping: dict[str, str] = {}
    rows = []
    for item in index["files"]:
        target = (workspace / _safe_destination(item["rendered"])).resolve()
        if not target.is_relative_to(workspace.resolve()):
            raise DRBenchDataError("Workspace symlink escapes the task directory.")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(derived_dir / item["rendered"], target)
        mapping[item["rendered"]] = item["destination"]
        rows.append(
            f"| `{item['destination']}` | {_APP_LABELS.get(item['app'], item['app'])} "
            f"| `{item['rendered']}` | {item['chars']:,} |"
        )
    (workspace / FILE_INDEX_NAME).write_text(
        "# Enterprise files available for this task\n\n"
        "Each file is provided as a text rendering. Read the rendering, but cite "
        "the original file name (first column).\n\n"
        "| Original file | Source | Text rendering | Characters |\n"
        "|---|---|---|---|\n" + "\n".join(rows) + "\n",
        encoding="utf-8",
    )
    leaked = [
        path for path in workspace.rglob("*") if path.name.casefold() in _FORBIDDEN_WORKSPACE_NAMES
    ]
    if leaked:
        raise DRBenchDataError(
            f"Refusing to run DRBench {task.task_id}: answer files in workspace: {leaked}"
        )
    return mapping


# ---------------------------------------------------------------------------
# Deterministic recall
# ---------------------------------------------------------------------------


def deterministic_recall(
    task: BenchmarkTask,
    report: str,
    *,
    tool_trace: list[dict[str, Any]],
    citation_keys: dict[str, list[str]] | None = None,
) -> dict[str, Any]:
    """Label-based file recall that needs no LLM judge.

    * ``file_read``: the Agent opened the file's rendering with ``read``.
    * ``file_surfaced``: read, or the file had a content match in a ``grep``
      result (``ls``/``find`` listings name every file, so they do not count).
    * ``cited``: the report names the file (its name or, for mailboxes and chat
      exports, an email subject or channel name contained in it).
    """

    files = task.reference.get("files") or []
    keys_by_destination = citation_keys or {}
    insight = {item["destination"] for item in files if item.get("qa_type") == "insight"}
    distractor = {item["destination"] for item in files if item.get("qa_type") == "distractor"}
    labeled = insight | distractor
    apps = {item["destination"]: item.get("app", "") for item in files}

    read_paths = {
        _destination_for_path(str(event.get("args", {}).get("path") or ""), labeled)
        for event in tool_trace
        if event.get("tool") == "read"
        and event.get("ok", True)
        and isinstance(event.get("args"), dict)
    }
    read = {value for value in read_paths if value}
    surfaced = set(read)
    for event in tool_trace:
        if event.get("tool") != "grep" or not event.get("ok", True):
            continue
        text = str(event.get("result") or "").casefold()
        for destination in labeled:
            if _mentions(text, PurePosixPath(destination).name.casefold()):
                surfaced.add(destination)

    folded_report = (report or "").casefold()
    cited = set()
    for destination in labeled:
        keys = keys_by_destination.get(destination) or _name_keys(PurePosixPath(destination).name)
        if any(key and _mentions(folded_report, key) for key in keys):
            cited.add(destination)

    def ratio(numerator: set[str], denominator: set[str]) -> float | None:
        return len(numerator & denominator) / len(denominator) if denominator else None

    by_app: dict[str, dict[str, Any]] = {}
    for app in sorted({apps[value] for value in insight}):
        app_insight = {value for value in insight if apps[value] == app}
        by_app[app] = {
            "insight_files": len(app_insight),
            "citation_recall": ratio(cited, app_insight),
            "file_read_recall": ratio(read, app_insight),
        }
    cited_labeled = cited & labeled
    qa = task.reference.get("qa") or []
    return {
        "evaluation_method": "drbench_deterministic_recall",
        "insight_files": len(insight),
        "distractor_files": len(distractor),
        "file_read_recall": ratio(read, insight),
        "file_surfaced_recall": ratio(surfaced, insight),
        "distractor_read_rate": ratio(read, distractor),
        "citation_recall": ratio(cited, insight),
        "citation_precision": (
            len(cited & insight) / len(cited_labeled) if cited_labeled else None
        ),
        "distractor_citation_rate": ratio(cited, distractor),
        "by_app": by_app,
        "external_insights": sum(
            item.get("qa_type") == "insight" and item.get("type") == "external_fact" for item in qa
        ),
        "files_read": sorted(read),
        "files_cited": sorted(cited),
        "insight_files_missed": sorted(insight - cited),
        "distractor_files_cited": sorted(cited & distractor),
    }


def _mentions(text: str, key: str) -> bool:
    """``key`` occurs in ``text`` as a whole token, so ``report.pdf`` does not
    match inside ``q3-report.pdf``. A following ``.`` is allowed (renderings
    are named ``<file>.md``; sentences end with periods)."""

    return re.search(rf"(?<![\w\-]){re.escape(key)}(?![\w\-])", text) is not None


def _destination_for_path(path: str, labeled: set[str]) -> str:
    name = PurePosixPath(path.replace("\\", "/")).name
    if name.endswith(RENDER_SUFFIX):
        name = name[: -len(RENDER_SUFFIX)]
    folded = name.casefold()
    for destination in labeled:
        if PurePosixPath(destination).name.casefold() == folded:
            return destination
    return ""


def citation_keys_for(task: BenchmarkTask) -> dict[str, list[str]]:
    index = ensure_rendered(task)
    return {item["destination"]: item.get("citation_keys") or [] for item in index["files"]}


__all__ = [
    "NAME",
    "DRBenchDataError",
    "citation_keys_for",
    "deterministic_recall",
    "ensure_rendered",
    "find_benchmark_dir",
    "load_task",
    "load_tasks",
    "render_file",
    "resolve_data_dir",
    "stage_workspace",
]
